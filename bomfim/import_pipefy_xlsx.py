"""Import registrations from Pipefy XLSX export."""
from __future__ import annotations

import re
from datetime import UTC, date, datetime
from pathlib import Path
from urllib.parse import unquote, urlparse

from openpyxl import load_workbook

from bomfim.default_owner import ensure_default_owner_user
from bomfim.extensions import db
from bomfim.import_clientes import UNIT_BY_CODE, ensure_reference_data
from bomfim.import_format import format_cep_masked, format_cnpj_masked, format_phone_display
from bomfim.anexo_download import download_anexos_for_registration
from bomfim.models import (
    CONTATO_TIPO_ENVIO_FATURAS,
    CONTATO_TIPO_OCORRENCIAS_COMERCIAL,
    DOWNLOAD_STATUS_PENDING,
    Contato,
    Organization,
    PotencialFaturamentoMes,
    Registration,
    RegistrationAnexo,
    RegistrationStage,
    TabelaPrecos,
    TipoCliente,
    TipoEmbalagem,
    TipoOportunidade,
    Unit,
    User,
)
from bomfim.security import hash_password

DEFAULT_IMPORT_PASSWORD = "123456789"

UNIT_NAME_ALIASES: dict[str, str] = {
    "RECIFE": "REC",
    "FORTALEZA": "FOR",
    "FEIRA DE SANTANA": "FST",
    "SALVADOR": "SSA",
    "MACEIO": "MCZ",
    "MACEIÓ": "MCZ",
    "ARACAJU": "AJU",
    "JOAO PESSOA": "JPA",
    "JOÃO PESSOA": "JPA",
    "NATAL": "NAT",
}

FASE_TO_SORT_ORDER: dict[str, int] = {
    "caixa de entrada": 0,
    "cadastro em andamento": 1,
    "aguardando assinatura": 2,
    "cadastro finalizado": 3,
    "cadastros declinados": 4,
}

FASE_METRIC_COLUMNS: dict[str, tuple[str, str, str]] = {
    "cadastro_em_andamento": (
        "Primeira vez que entrou na fase Cadastro em Andamento",
        "Última vez que saiu da fase Cadastro em Andamento",
        "Tempo total na fase Cadastro em Andamento (dias)",
    ),
    "caixa_entrada": (
        "Primeira vez que entrou na fase Caixa de entrada",
        "Última vez que saiu da fase Caixa de entrada",
        "Tempo total na fase Caixa de entrada (dias)",
    ),
    "cadastros_declinados": (
        "Primeira vez que entrou na fase Cadastros Declinados",
        "Última vez que saiu da fase Cadastros Declinados",
        "Tempo total na fase Cadastros Declinados (dias)",
    ),
    "aguardando_assinatura": (
        "Primeira vez que entrou na fase Aguardando Assinatura",
        "Última vez que saiu da fase Aguardando Assinatura",
        "Tempo total na fase Aguardando Assinatura (dias)",
    ),
    "cadastro_finalizado": (
        "Primeira vez que entrou na fase Cadastro Finalizado",
        "Última vez que saiu da fase Cadastro Finalizado",
        "Tempo total na fase Cadastro Finalizado (dias)",
    ),
}

_ANEXO_URL_SPLIT = re.compile(r",\s*(?=https?://)", re.IGNORECASE)


def _norm_header(value: object) -> str:
    return " ".join(str(value or "").split())


def _cell_str(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return "True" if value else "False"
    text = str(value).strip()
    return text or None


def _parse_datetime(value: object) -> datetime | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.replace(tzinfo=UTC) if value.tzinfo is None else value
    text = str(value).strip()
    if not text:
        return None
    for fmt in (
        "%Y-%m-%d %H:%M:%S.%f",
        "%Y-%m-%d %H:%M:%S",
        "%d/%m/%Y %H:%M:%S",
        "%d/%m/%Y",
    ):
        try:
            dt = datetime.strptime(text[:26], fmt)
            return dt.replace(tzinfo=UTC)
        except ValueError:
            continue
    return None


def _parse_date_from_cell(value: object) -> date:
    dt = _parse_datetime(value)
    if dt:
        return dt.date()
    return date.today()


def _catalog_key(text: str) -> str:
    return " ".join(text.upper().split())


def _build_catalog_cache() -> dict[str, dict[str, int]]:
    caches: dict[str, dict[str, int]] = {
        "tipo_oportunidade": {},
        "tabela_precos": {},
        "tipo_cliente": {},
        "potencial_faturamento_mes": {},
        "tipo_embalagem": {},
    }
    for row in TipoOportunidade.query.all():
        caches["tipo_oportunidade"][_catalog_key(row.descricao)] = row.id
    for row in TabelaPrecos.query.all():
        caches["tabela_precos"][_catalog_key(row.descricao)] = row.id
    for row in TipoCliente.query.all():
        caches["tipo_cliente"][_catalog_key(row.descricao)] = row.id
    for row in PotencialFaturamentoMes.query.all():
        caches["potencial_faturamento_mes"][_catalog_key(row.descricao)] = row.id
    for row in TipoEmbalagem.query.all():
        caches["tipo_embalagem"][_catalog_key(row.descricao)] = row.id
    return caches


def _resolve_catalog_id(cache: dict[str, int], raw: str | None) -> int | None:
    if not raw:
        return None
    key = _catalog_key(raw)
    if key in cache:
        return cache[key]
    for desc_key, pk in cache.items():
        if key in desc_key or desc_key in key:
            return pk
    return None


def _resolve_stage_id(fase: str | None, warnings: list[str]) -> int | None:
    if not fase:
        return _stage_by_sort_order(1)
    norm = fase.strip().lower()
    sort_order = FASE_TO_SORT_ORDER.get(norm)
    if sort_order is None:
        warnings.append(f"Fase desconhecida {fase!r}; usando Cadastro em Andamento.")
        sort_order = 1
    return _stage_by_sort_order(sort_order)


def _stage_by_sort_order(sort_order: int) -> int | None:
    stage = RegistrationStage.query.filter_by(sort_order=sort_order).first()
    return stage.id if stage else None


def _resolve_unit(org: Organization, raw_unit: str | None) -> Unit:
    label = _cell_str(raw_unit) or "Salvador"
    key = label.upper()
    code = UNIT_NAME_ALIASES.get(key, label[:16].upper())
    unit = Unit.query.filter((Unit.code == code) | (Unit.name.ilike(label))).first()
    if unit:
        return unit
    name, locality = UNIT_BY_CODE.get(code, (label, label))
    unit = Unit(
        organization_id=org.id,
        name=name,
        code=code,
        locality=locality,
        responsible_name="—",
        status="active",
    )
    db.session.add(unit)
    db.session.flush()
    return unit


def _promoter_user(org_id: int, name: str | None, email: str | None, cache: dict[str, User]) -> User | None:
    mail = (_cell_str(email) or "").lower()
    if not mail:
        return None
    if mail in cache:
        return cache[mail]
    user = User.query.filter_by(email=mail).first()
    if user is None:
        user = User(
            organization_id=org_id,
            name=_cell_str(name) or mail.split("@")[0],
            email=mail,
            role="promoter",
            password_hash=hash_password(DEFAULT_IMPORT_PASSWORD),
            status="active",
        )
        db.session.add(user)
        db.session.flush()
    cache[mail] = user
    return user


def _row_dict(headers: list[str], row: tuple) -> dict[str, object | None]:
    data: dict[str, object | None] = {}
    for idx, header in enumerate(headers):
        if not header:
            continue
        data[header] = row[idx] if idx < len(row) else None
    return data


def _parse_anexo_urls(raw: str | None) -> list[str]:
    text = _cell_str(raw)
    if not text:
        return []
    parts = _ANEXO_URL_SPLIT.split(text)
    urls: list[str] = []
    for part in parts:
        url = part.strip().rstrip(",")
        if url.lower().startswith("http"):
            urls.append(url)
    return urls


def _filename_from_url(url: str) -> str:
    path = unquote(urlparse(url).path)
    name = path.rsplit("/", 1)[-1]
    return name or "anexo"


def _fase_metricas(row: dict[str, object | None]) -> dict | None:
    payload: dict[str, dict[str, object | None]] = {}
    for key, (enter_col, leave_col, days_col) in FASE_METRIC_COLUMNS.items():
        payload[key] = {
            "primeira_entrada": _serialize_dt(row.get(enter_col)),
            "ultima_saida": _serialize_dt(row.get(leave_col)),
            "tempo_total_dias": row.get(days_col),
        }
    return payload


def _serialize_dt(value: object | None) -> str | None:
    dt = _parse_datetime(value)
    return dt.isoformat() if dt else _cell_str(value)


def _contact_has_data(**fields: str | None) -> bool:
    return any(fields.values())


def _replace_contatos(reg: Registration, row: dict[str, object | None]) -> int:
    with db.session.no_autoflush:
        Contato.query.filter(
            Contato.registration_id == reg.id,
            Contato.tipo_contato.in_([CONTATO_TIPO_ENVIO_FATURAS, CONTATO_TIPO_OCORRENCIAS_COMERCIAL]),
        ).delete(synchronize_session=False)

    created = 0
    fat_email = _cell_str(row.get("E-mail's para envio Faturas/DACTE (XML)"))
    fat_nome = _cell_str(row.get("Nome para envio de faturas"))
    fat_tel = format_phone_display(_cell_str(row.get("Telefone para envio de faturas")))
    if _contact_has_data(nome=fat_nome, email=fat_email, telefone1=fat_tel):
        db.session.add(
            Contato(
                registration_id=reg.id,
                tipo_contato=CONTATO_TIPO_ENVIO_FATURAS,
                nome=fat_nome or fat_email or "Envio faturas",
                email=fat_email,
                telefone1=fat_tel,
                status=True,
            )
        )
        created += 1

    ocorr_email = _cell_str(row.get("Ocorrências administrativas (E-mail)"))
    com_nome = _cell_str(row.get("Comercial e Marketing (Nome)"))
    com_email = _cell_str(row.get("Comercial e Marketing (E-mail):"))
    com_tel = format_phone_display(_cell_str(row.get("Comercial e Marketing (Telefone)")))
    if _contact_has_data(
        nome=com_nome,
        email=com_email,
        telefone1=com_tel,
        email_ocorrencias=ocorr_email,
    ):
        db.session.add(
            Contato(
                registration_id=reg.id,
                tipo_contato=CONTATO_TIPO_OCORRENCIAS_COMERCIAL,
                nome=com_nome or com_email or "Comercial",
                email=com_email,
                telefone1=com_tel,
                email_ocorrencias=ocorr_email,
                status=True,
            )
        )
        created += 1
    return created


def _replace_anexos(reg: Registration, row: dict[str, object | None]) -> int:
    with db.session.no_autoflush:
        RegistrationAnexo.query.filter_by(registration_id=reg.id).delete(synchronize_session=False)
    urls = _parse_anexo_urls(_cell_str(row.get("Anexos (Tabela de Preços, Documentos e Licenças)")))
    for ordem, url in enumerate(urls):
        db.session.add(
            RegistrationAnexo(
                registration_id=reg.id,
                url=url,
                filename=_filename_from_url(url),
                ordem=ordem,
                download_status=DOWNLOAD_STATUS_PENDING,
            )
        )
    return len(urls)


def _registration_payload(
    org: Organization,
    row: dict[str, object | None],
    *,
    catalogs: dict[str, dict[str, int]],
    default_owner: User,
    promoter_cache: dict[str, User],
    warnings: list[str],
) -> dict:
    cnpj = format_cnpj_masked(_cell_str(row.get("CNPJ:")))
    if not cnpj:
        raise ValueError("CNPJ ausente")

    razao = _cell_str(row.get("Razão Social"))
    fantasia = _cell_str(row.get("Nome Fantasia:"))
    titulo = _cell_str(row.get("Título"))
    name = razao or fantasia or titulo or cnpj

    cidade = _cell_str(row.get("Cidade::")) or _cell_str(row.get("Cidade:")) or ""
    uf = _cell_str(row.get("Estado::")) or _cell_str(row.get("Estado:")) or ""
    city_label = f"{cidade}, {uf}".strip(", ") if uf else cidade

    tabela_txt = _cell_str(row.get("Tabela de Preços")) or "—"
    unit = _resolve_unit(org, _cell_str(row.get("Unidade:")))
    promoter = _promoter_user(
        org.id,
        _cell_str(row.get("Promotor de Vendas (Nome):")),
        _cell_str(row.get("Promotor de Vendas (E-mail)")),
        promoter_cache,
    )

    com_email = _cell_str(row.get("Comercial e Marketing (E-mail):"))
    stage_id = _resolve_stage_id(_cell_str(row.get("Fase atual")), warnings)

    return dict(
        organization_id=org.id,
        name=name,
        cnpj=cnpj,
        contact_email=com_email,
        contact_name=_cell_str(row.get("Comercial e Marketing (Nome)")),
        postal_code=format_cep_masked(_cell_str(row.get("Cep:")) or _cell_str(row.get("CEP:"))),
        price_table=tabela_txt,
        stage_id=stage_id,
        city=city_label or unit.locality,
        unit_id=unit.id,
        promoter_id=promoter.id if promoter else default_owner.id,
        owner_id=promoter.id if promoter else default_owner.id,
        updated_at=_parse_date_from_cell(row.get("Atualizado em")),
        pipefy_codigo=_cell_str(row.get("Código")),
        titulo=titulo,
        moskit_id=_cell_str(row.get("ID Moskit")),
        criado_em=_parse_datetime(row.get("Criado em")),
        atualizado_em=_parse_datetime(row.get("Atualizado em")),
        finalizado_em=_parse_datetime(row.get("Finalizado em")),
        tipo_oportunidade_id=_resolve_catalog_id(
            catalogs["tipo_oportunidade"], _cell_str(row.get("Tipo da Oportunidade"))
        ),
        tabela_precos_id=_resolve_catalog_id(catalogs["tabela_precos"], tabela_txt),
        tipo_cliente_id=_resolve_catalog_id(catalogs["tipo_cliente"], _cell_str(row.get("Tipo de Cliente"))),
        potencial_faturamento_mes_id=_resolve_catalog_id(
            catalogs["potencial_faturamento_mes"],
            _cell_str(row.get("Potencial de Faturamento por Mês (R$)")),
        ),
        tipo_embalagem_id=_resolve_catalog_id(
            catalogs["tipo_embalagem"], _cell_str(row.get("Tipo da Embalagem"))
        ),
        razao_social=razao,
        nome_fantasia=fantasia,
        segmento=_cell_str(row.get("Segmento")),
        natureza_carga=_cell_str(row.get("Natureza de Carga")),
        inscricao_estadual=_cell_str(row.get("Inscrição Estadual")),
        substituto_tributario_iss=_cell_str(row.get("É substituto tributário de ISS?")),
        inscricao_municipal=_cell_str(row.get("Inscrição Municipal")),
        endereco=_cell_str(row.get("Endereço:")),
        numero=_cell_str(row.get("Número")),
        complemento=_cell_str(row.get("Complemento:")),
        bairro=_cell_str(row.get("Bairro::")) or _cell_str(row.get("Bairro:")),
        estado=uf or None,
        igual_endereco_cliente=_cell_str(row.get("Igual ao endereço do cliente?")) == "True",
        endereco_cobranca=_cell_str(row.get("Endereco")),
        numero_cobranca=_cell_str(row.get("Numero")),
        complemento_cobranca=_cell_str(row.get("Comp.")),
        cep_cobranca=format_cep_masked(_cell_str(row.get("CEP:"))),
        bairro_cobranca=_cell_str(row.get("Bairro:")),
        cidade_cobranca=_cell_str(row.get("Cidade:")),
        estado_cobranca=_cell_str(row.get("Estado:")),
        promotor_nome=_cell_str(row.get("Promotor de Vendas (Nome):")),
        promotor_email=_cell_str(row.get("Promotor de Vendas (E-mail)")),
        operador_logistico=_cell_str(row.get("Se Operador Logístico")),
        contrato_proprio=_cell_str(row.get("Cliente tem contrato próprio?")),
        licenca_transporte=_cell_str(row.get("Alguma licença de transporte específica para a operação?")),
        procedimento_especifico=_cell_str(row.get("Algum procedimento específico:")),
        responsabilidade_seguro=_cell_str(row.get("Responsabilidade do Seguro")),
        procedimento_devolucao_comprovantes=_cell_str(
            row.get("Existe algum procedimento específico para devolução de comprovantes de entrega?")
        ),
        detalhamento_envio=_cell_str(row.get("Detalhamento do envio")),
        endereco_envio_fisico=_cell_str(row.get("Endereço para envio físico")),
        email_envio_digital=_cell_str(row.get("E-mail para envio digital")),
        edi_trabalha=_cell_str(
            row.get(
                "O cliente trabalha com EDI (Eletronic Date Interchange – Troca Eletrônica de Dados):"
            )
        ),
        edi_tipos_arquivo=_cell_str(row.get("Quais os tipos de arquivo que serão solicitados?")),
        edi_frequencia=_cell_str(row.get("Qual a frequência de envio dos arquivos solicitados?")),
        edi_horarios=_cell_str(row.get("Horários:")),
        edi_contato=_cell_str(row.get("Informar contato:")),
        envio_notfis=_cell_str(row.get("Existe envio de NOTFIS?")),
        observacoes=_cell_str(row.get("Observações")),
        coleta_automatica=_cell_str(row.get("Coleta Automática?")),
        necessita_assinatura_digital=_cell_str(row.get("Necessita de assinatura digital?")),
        nome_assinatura=_cell_str(row.get("Nome para assinatura")),
        email_assinatura=_cell_str(row.get("E-mail para assinatura")),
        fase_metricas=_fase_metricas(row),
    )


def import_pipefy_xlsx(path: Path, *, batch_size: int = 50, download_anexos: bool = False) -> dict[str, int | list[str]]:
    org = ensure_reference_data()
    default_owner = ensure_default_owner_user(org)
    catalogs = _build_catalog_cache()
    promoter_cache: dict[str, User] = {}
    warnings: list[str] = []

    wb = load_workbook(path, read_only=True, data_only=True)
    sheet = wb["Report"] if "Report" in wb.sheetnames else wb.active
    rows_iter = sheet.iter_rows(values_only=True)
    header_row = next(rows_iter)
    headers = [_norm_header(h) for h in header_row]

    stats: dict[str, int | list[str]] = {
        "created": 0,
        "updated": 0,
        "skipped": 0,
        "contacts": 0,
        "anexos": 0,
        "anexos_downloaded_ok": 0,
        "anexos_download_failed": 0,
        "warnings": warnings,
    }

    pending = 0
    for row in rows_iter:
        data = _row_dict(headers, row)
        try:
            payload = _registration_payload(
                org,
                data,
                catalogs=catalogs,
                default_owner=default_owner,
                promoter_cache=promoter_cache,
                warnings=warnings,
            )
        except ValueError:
            stats["skipped"] = int(stats["skipped"]) + 1
            continue

        cnpj = payload["cnpj"]
        reg = Registration.query.filter_by(cnpj=cnpj).first()
        anexo_urls = _parse_anexo_urls(
            _cell_str(data.get("Anexos (Tabela de Preços, Documentos e Licenças)"))
        )
        payload["documents_count"] = len(anexo_urls)

        if reg is None:
            reg = Registration(**payload)
            db.session.add(reg)
            db.session.flush()
            stats["created"] = int(stats["created"]) + 1
        else:
            for key, value in payload.items():
                setattr(reg, key, value)
            stats["updated"] = int(stats["updated"]) + 1

        stats["contacts"] = int(stats["contacts"]) + _replace_contatos(reg, data)
        stats["anexos"] = int(stats["anexos"]) + _replace_anexos(reg, data)
        if download_anexos:
            db.session.flush()
            dl = download_anexos_for_registration(reg.id)
            stats["anexos_downloaded_ok"] = int(stats["anexos_downloaded_ok"]) + dl["ok"]
            stats["anexos_download_failed"] = int(stats["anexos_download_failed"]) + dl["failed"]

        pending += 1
        if pending >= batch_size:
            db.session.commit()
            pending = 0

    db.session.commit()
    wb.close()
    return stats
