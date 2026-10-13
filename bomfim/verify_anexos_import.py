from __future__ import annotations

import csv
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

from bomfim.extensions import db
from bomfim.import_format import format_cnpj_masked
from bomfim.import_pipefy_xlsx import _norm_header, _parse_anexo_urls, _row_dict
from bomfim.models import DOWNLOAD_STATUS_OK, Registration, RegistrationAnexo
from bomfim.storage_paths import anexo_absolute_path, get_storage_root


def _short_url(url: str, max_len: int = 80) -> str:
    return url if len(url) <= max_len else url[: max_len - 3] + "..."


def verify_anexos_import(*, xlsx_path: Path | None = None, output_csv: Path | None = None) -> dict:
    summary: dict[str, int | list] = {
        "total_db": 0,
        "ok_with_file": 0,
        "pending": 0,
        "failed": 0,
        "ok_missing_file": 0,
        "orphan_files": 0,
        "missing_in_db": 0,
        "extra_in_db": 0,
        "issues": [],
    }

    rows = RegistrationAnexo.query.order_by(RegistrationAnexo.id).all()
    summary["total_db"] = len(rows)

    known_paths: set[str] = set()
    issues: list[dict[str, str]] = []

    for row in rows:
        st = row.download_status
        if st == "pending":
            summary["pending"] = int(summary["pending"]) + 1
            issues.append(
                {
                    "kind": "pending",
                    "anexo_id": str(row.id),
                    "registration_id": str(row.registration_id),
                    "filename": row.filename or "",
                    "detail": _short_url(row.url),
                }
            )
        elif st == "failed":
            summary["failed"] = int(summary["failed"]) + 1
            issues.append(
                {
                    "kind": "failed",
                    "anexo_id": str(row.id),
                    "registration_id": str(row.registration_id),
                    "filename": row.filename or "",
                    "detail": (row.download_error or "")[:500],
                }
            )
        elif st == DOWNLOAD_STATUS_OK:
            if row.local_path and anexo_absolute_path(row.local_path).is_file():
                summary["ok_with_file"] = int(summary["ok_with_file"]) + 1
                known_paths.add(row.local_path.replace("\\", "/"))
            else:
                summary["ok_missing_file"] = int(summary["ok_missing_file"]) + 1
                issues.append(
                    {
                        "kind": "ok_missing_file",
                        "anexo_id": str(row.id),
                        "registration_id": str(row.registration_id),
                        "filename": row.filename or "",
                        "detail": row.local_path or "",
                    }
                )

    root = get_storage_root() / "registration_anexos"
    orphan_count = 0
    if root.is_dir():
        for path in root.rglob("*"):
            if not path.is_file() or path.name.endswith(".part"):
                continue
            rel = path.relative_to(get_storage_root()).as_posix()
            if rel not in known_paths:
                orphan_count += 1
                issues.append(
                    {
                        "kind": "orphan_file",
                        "anexo_id": "",
                        "registration_id": "",
                        "filename": path.name,
                        "detail": rel,
                    }
                )
    summary["orphan_files"] = orphan_count

    if xlsx_path is not None:
        xlsx_urls_by_cnpj: dict[str, set[str]] = defaultdict(set)
        wb = load_workbook(xlsx_path, read_only=True, data_only=True)
        sheet = wb["Report"] if "Report" in wb.sheetnames else wb.active
        headers = [_norm_header(h) for h in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
        for row in sheet.iter_rows(min_row=2, values_only=True):
            data = _row_dict(headers, row)
            cnpj = format_cnpj_masked(str(data.get("CNPJ:") or ""))
            if not cnpj:
                continue
            for url in _parse_anexo_urls(str(data.get("Anexos (Tabela de Preços, Documentos e Licenças)") or "")):
                xlsx_urls_by_cnpj[cnpj].add(url)
        wb.close()

        db_urls_by_cnpj: dict[str, set[str]] = defaultdict(set)
        for reg in Registration.query.all():
            for anexo in reg.anexos:
                db_urls_by_cnpj[reg.cnpj].add(anexo.url)

        all_cnpjs = set(xlsx_urls_by_cnpj) | set(db_urls_by_cnpj)
        missing = extra = 0
        for cnpj in all_cnpjs:
            in_xlsx = xlsx_urls_by_cnpj.get(cnpj, set())
            in_db = db_urls_by_cnpj.get(cnpj, set())
            for url in in_xlsx - in_db:
                missing += 1
                issues.append(
                    {
                        "kind": "missing_in_db",
                        "anexo_id": "",
                        "registration_id": cnpj,
                        "filename": "",
                        "detail": _short_url(url),
                    }
                )
            for url in in_db - in_xlsx:
                extra += 1
                issues.append(
                    {
                        "kind": "extra_in_db",
                        "anexo_id": "",
                        "registration_id": cnpj,
                        "filename": "",
                        "detail": _short_url(url),
                    }
                )
        summary["missing_in_db"] = missing
        summary["extra_in_db"] = extra

    summary["issues"] = issues

    if output_csv is not None:
        output_csv.parent.mkdir(parents=True, exist_ok=True)
        with output_csv.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(
                handle, fieldnames=["kind", "anexo_id", "registration_id", "filename", "detail"]
            )
            writer.writeheader()
            writer.writerows(issues)

    return summary
